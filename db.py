"""Database/repository layer. Keep SQL here so the UI can later migrate to an API."""
import os
import sqlite3
import sys
from datetime import datetime, date

APP_DIR = (
    os.path.dirname(os.path.abspath(sys.executable))
    if getattr(sys, "frozen", False)
    else os.path.dirname(os.path.abspath(__file__))
)
DATA_DIR = os.path.join(APP_DIR, "data")
DB_PATH = os.path.join(DATA_DIR, "asset_management.db")

ASSET_PREFIXES = {
    "Desktop CPU": "CPU",
    "Monitor": "MON",
    "Keyboard": "KBD",
    "Mouse": "MOU",
    "Printer": "PRN",
    "Server": "SER",
    "UPS": "UPS",
    "Scanner": "SCN",
    "Projector": "PRO",
    "Laptop": "LAP",
    "Other": "OTH",
}

def connect():
    os.makedirs(DATA_DIR, exist_ok=True)
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    return con

def init_db():
    with connect() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS departments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT NOT NULL UNIQUE,
            name TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS locations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            department_id INTEGER,
            block TEXT DEFAULT '',
            floor TEXT DEFAULT '',
            room TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY(department_id) REFERENCES departments(id)
        );
        CREATE TABLE IF NOT EXISTS asset_groups (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            parent_code TEXT NOT NULL UNIQUE,
            group_name TEXT NOT NULL,
            group_type TEXT NOT NULL DEFAULT 'Room',
            department_id INTEGER,
            location_id INTEGER,
            staff_name TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY(department_id) REFERENCES departments(id),
            FOREIGN KEY(location_id) REFERENCES locations(id)
        );
        CREATE TABLE IF NOT EXISTS computer_setups (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            setup_code TEXT NOT NULL UNIQUE,
            setup_name TEXT NOT NULL,
            group_id INTEGER,
            notes TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY(group_id) REFERENCES asset_groups(id)
        );
        CREATE TABLE IF NOT EXISTS assets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            asset_code TEXT NOT NULL UNIQUE,
            device_type TEXT NOT NULL,
            brand TEXT DEFAULT '',
            model TEXT DEFAULT '',
            manufacturer_serial TEXT DEFAULT '',
            processor TEXT DEFAULT '',
            ram TEXT DEFAULT '',
            storage TEXT DEFAULT '',
            operating_system TEXT DEFAULT '',
            department_id INTEGER,
            location_id INTEGER,
            group_id INTEGER,
            setup_id INTEGER,
            staff_name TEXT DEFAULT '',
            status TEXT NOT NULL DEFAULT 'Working',
            purchase_date TEXT DEFAULT '',
            warranty_expiry TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY(department_id) REFERENCES departments(id),
            FOREIGN KEY(location_id) REFERENCES locations(id),
            FOREIGN KEY(group_id) REFERENCES asset_groups(id),
            FOREIGN KEY(setup_id) REFERENCES computer_setups(id)
        );
        CREATE TABLE IF NOT EXISTS fault_reports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fault_code TEXT NOT NULL UNIQUE,
            asset_id INTEGER NOT NULL,
            reported_date TEXT NOT NULL,
            description TEXT NOT NULL,
            reported_by TEXT DEFAULT '',
            priority TEXT DEFAULT 'Normal',
            status TEXT NOT NULL DEFAULT 'Open',
            resolution TEXT DEFAULT '',
            resolved_date TEXT DEFAULT '',
            FOREIGN KEY(asset_id) REFERENCES assets(id)
        );
        CREATE TABLE IF NOT EXISTS repair_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            asset_id INTEGER NOT NULL,
            fault_id INTEGER,
            repair_date TEXT NOT NULL,
            technician TEXT DEFAULT '',
            action_taken TEXT DEFAULT '',
            cost REAL DEFAULT 0,
            notes TEXT DEFAULT '',
            FOREIGN KEY(asset_id) REFERENCES assets(id),
            FOREIGN KEY(fault_id) REFERENCES fault_reports(id)
        );
        CREATE TABLE IF NOT EXISTS asset_movements (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            asset_id INTEGER NOT NULL,
            movement_date TEXT NOT NULL,
            from_location TEXT DEFAULT '',
            to_location TEXT NOT NULL,
            reason TEXT DEFAULT '',
            moved_by TEXT DEFAULT '',
            department_id INTEGER,
            group_id INTEGER,
            FOREIGN KEY(asset_id) REFERENCES assets(id)
        );
        CREATE TABLE IF NOT EXISTS setup_assets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            setup_id INTEGER NOT NULL,
            asset_id INTEGER NOT NULL,
            start_date TEXT NOT NULL,
            end_date TEXT DEFAULT '',
            FOREIGN KEY(setup_id) REFERENCES computer_setups(id),
            FOREIGN KEY(asset_id) REFERENCES assets(id)
        );
        CREATE INDEX IF NOT EXISTS idx_assets_code ON assets(asset_code);
        CREATE INDEX IF NOT EXISTS idx_fault_asset ON fault_reports(asset_id);
        """)
        movement_columns = {
            row["name"] for row in con.execute("PRAGMA table_info(asset_movements)")
        }
        if "department_id" not in movement_columns:
            con.execute("ALTER TABLE asset_movements ADD COLUMN department_id INTEGER")
        if "group_id" not in movement_columns:
            con.execute("ALTER TABLE asset_movements ADD COLUMN group_id INTEGER")

def _now():
    return datetime.now().isoformat(timespec="seconds")

def list_departments():
    with connect() as con:
        return con.execute("SELECT * FROM departments ORDER BY name").fetchall()

def add_department(code, name):
    code, name = code.strip().upper(), name.strip()
    if not code or not name:
        raise ValueError("Department code and name are required.")
    with connect() as con:
        con.execute("INSERT INTO departments(code,name,created_at) VALUES(?,?,?)",
                    (code, name, _now()))

def list_groups():
    with connect() as con:
        return con.execute("""
            SELECT g.*, d.name AS department_name, l.block, l.floor, l.room
            FROM asset_groups g
            LEFT JOIN departments d ON d.id=g.department_id
            LEFT JOIN locations l ON l.id=g.location_id
            ORDER BY g.group_name
        """).fetchall()

def list_group_assets(group_id):
    with connect() as con:
        return con.execute("""
            SELECT DISTINCT a.*, s.setup_code, s.setup_name
            FROM assets a
            LEFT JOIN computer_setups s ON s.id=a.setup_id
            WHERE a.group_id=?
            ORDER BY a.device_type, a.asset_code
        """, (group_id,)).fetchall()

def create_group(group_name, group_type, department_id, block, floor, room, staff_name, notes):
    group_name = group_name.strip()
    if not group_name:
        raise ValueError("Group name is required.")
    code_seed = "".join(ch for ch in group_name.upper() if ch.isalnum())[:3] or "GRP"
    with connect() as con:
        cur = con.cursor()
        location_id = None
        if department_id:
            cur.execute("""INSERT INTO locations(department_id,block,floor,room,created_at)
                           VALUES(?,?,?,?,?)""",
                        (department_id, block.strip(), floor.strip(), room.strip(), _now()))
            location_id = cur.lastrowid
        # Allocate next unique parent code, independent of the room's text.
        row = cur.execute("SELECT COUNT(*) AS n FROM asset_groups WHERE parent_code LIKE ?",
                          (f"APS-{code_seed}-%",)).fetchone()
        n = row["n"] + 1
        parent_code = f"APS-{code_seed}-{n:03d}"
        while cur.execute("SELECT 1 FROM asset_groups WHERE parent_code=?", (parent_code,)).fetchone():
            n += 1
            parent_code = f"APS-{code_seed}-{n:03d}"
        cur.execute("""INSERT INTO asset_groups(parent_code,group_name,group_type,department_id,
                       location_id,staff_name,notes,created_at) VALUES(?,?,?,?,?,?,?,?)""",
                    (parent_code, group_name, group_type, department_id, location_id,
                     staff_name.strip(), notes.strip(), _now()))
        return parent_code

def list_setups():
    with connect() as con:
        return con.execute("""
          SELECT s.*, g.parent_code, g.group_name
          FROM computer_setups s LEFT JOIN asset_groups g ON g.id=s.group_id
          ORDER BY s.setup_code
        """).fetchall()

def create_setup(setup_name, group_id, notes=""):
    with connect() as con:
        cur = con.cursor()
        row = cur.execute("SELECT COUNT(*) AS n FROM computer_setups").fetchone()
        n = row["n"] + 1
        code = f"SET-{n:04d}"
        while cur.execute("SELECT 1 FROM computer_setups WHERE setup_code=?", (code,)).fetchone():
            n += 1
            code = f"SET-{n:04d}"
        cur.execute("""INSERT INTO computer_setups(setup_code,setup_name,group_id,notes,created_at)
                       VALUES(?,?,?,?,?)""", (code, setup_name.strip(), group_id, notes.strip(), _now()))
        return code

def list_assets(search=""):
    search = f"%{search.strip()}%"
    with connect() as con:
        return con.execute("""
          SELECT a.*, d.name AS department_name, l.block, l.floor, l.room,
                 g.parent_code, g.group_name, s.setup_code, s.setup_name
          FROM assets a
          LEFT JOIN departments d ON d.id=a.department_id
          LEFT JOIN locations l ON l.id=a.location_id
          LEFT JOIN asset_groups g ON g.id=a.group_id
          LEFT JOIN computer_setups s ON s.id=a.setup_id
          WHERE a.asset_code LIKE ? OR a.device_type LIKE ? OR a.brand LIKE ?
             OR a.model LIKE ? OR a.manufacturer_serial LIKE ? OR d.name LIKE ?
             OR g.group_name LIKE ? OR a.status LIKE ?
          ORDER BY a.id DESC
        """, (search, search, search, search, search, search, search, search)).fetchall()

def get_asset(asset_id):
    with connect() as con:
        return con.execute("SELECT * FROM assets WHERE id=?", (asset_id,)).fetchone()

def _next_asset_code(con, device_type):
    prefix = ASSET_PREFIXES.get(device_type, "OTH")
    row = con.execute("SELECT asset_code FROM assets WHERE asset_code LIKE ? ORDER BY id DESC LIMIT 1",
                      (f"APS-{prefix}-%",)).fetchone()
    number = 1
    if row:
        try:
            number = int(row["asset_code"].rsplit("-", 1)[1]) + 1
        except (ValueError, IndexError):
            number = con.execute("SELECT COUNT(*) FROM assets WHERE asset_code LIKE ?",
                                 (f"APS-{prefix}-%",)).fetchone()[0] + 1
    code = f"APS-{prefix}-{number:04d}"
    while con.execute("SELECT 1 FROM assets WHERE asset_code=?", (code,)).fetchone():
        number += 1
        code = f"APS-{prefix}-{number:04d}"
    return code

def save_asset(data, asset_id=None):
    required = ("device_type",)
    if not data.get("device_type"):
        raise ValueError("Device type is required.")
    now = _now()
    fields = ["device_type","brand","model","manufacturer_serial","processor","ram","storage",
              "operating_system","department_id","location_id","group_id","setup_id","staff_name",
              "status","purchase_date","warranty_expiry","notes"]
    vals = [data.get(k) for k in fields]
    with connect() as con:
        if asset_id:
            sets = ", ".join(f"{f}=?" for f in fields)
            con.execute(f"UPDATE assets SET {sets}, updated_at=? WHERE id=?",
                        vals + [now, asset_id])
            return get_asset(asset_id)["asset_code"]
        code = _next_asset_code(con, data["device_type"])
        placeholders = ",".join("?" for _ in fields)
        con.execute(f"""INSERT INTO assets(asset_code,{",".join(fields)},created_at,updated_at)
                        VALUES(?,{placeholders},?,?)""",
                    [code] + vals + [now, now])
        return code

def add_fault(asset_code, reported_date, description, reported_by, priority):
    with connect() as con:
        asset = con.execute("SELECT id FROM assets WHERE asset_code=?", (asset_code,)).fetchone()
        if not asset:
            raise ValueError("Asset ID was not found.")
        n = con.execute("SELECT COUNT(*) FROM fault_reports").fetchone()[0] + 1
        fault_code = f"FLT-{n:05d}"
        con.execute("""INSERT INTO fault_reports(fault_code,asset_id,reported_date,description,
                       reported_by,priority,status) VALUES(?,?,?,?,?,?, 'Open')""",
                    (fault_code, asset["id"], reported_date, description.strip(),
                     reported_by.strip(), priority))
        con.execute("UPDATE assets SET status='Faulty', updated_at=? WHERE id=?",
                    (_now(), asset["id"]))
        return fault_code

def list_faults(search=""):
    search = f"%{search.strip()}%"
    with connect() as con:
        return con.execute("""
          SELECT f.*, a.asset_code, a.device_type, a.brand
          FROM fault_reports f JOIN assets a ON a.id=f.asset_id
          WHERE f.fault_code LIKE ? OR a.asset_code LIKE ? OR f.description LIKE ?
             OR f.status LIKE ?
          ORDER BY f.id DESC
        """, (search, search, search, search)).fetchall()

def get_asset_history(asset_code):
    with connect() as con:
        asset = con.execute(
            "SELECT id, asset_code, device_type, brand, model, status FROM assets WHERE asset_code=?",
            (asset_code.strip(),),
        ).fetchone()
        if not asset:
            raise ValueError("Asset ID was not found.")
        return asset, con.execute("""
            SELECT event_date, event_type, reference, details, performed_by, status
            FROM (
                SELECT f.reported_date AS event_date, 'Fault reported' AS event_type,
                       f.fault_code AS reference, f.description AS details,
                       f.reported_by AS performed_by, f.status AS status, f.id AS event_id
                FROM fault_reports f
                WHERE f.asset_id=?
                UNION ALL
                SELECT r.repair_date, 'Repair recorded', f.fault_code,
                       r.action_taken, r.technician,
                       'Cost: ' || printf('%.2f', r.cost), r.id
                FROM repair_records r
                LEFT JOIN fault_reports f ON f.id=r.fault_id
                WHERE r.asset_id=?
                UNION ALL
                SELECT m.movement_date, 'Asset moved', 'MOV-' || m.id,
                       m.from_location || ' -> ' || m.to_location,
                       m.moved_by, 'Location change', m.id
                FROM asset_movements m
                WHERE m.asset_id=?
            )
            ORDER BY event_date DESC, event_id DESC
        """, (asset["id"], asset["id"], asset["id"])).fetchall()

def resolve_fault(fault_id, resolution, technician, repair_date, cost):
    with connect() as con:
        row = con.execute("SELECT asset_id FROM fault_reports WHERE id=?", (fault_id,)).fetchone()
        if not row:
            raise ValueError("Fault report not found.")
        con.execute("""UPDATE fault_reports SET status='Resolved', resolution=?, resolved_date=?
                       WHERE id=?""", (resolution.strip(), repair_date, fault_id))
        con.execute("""INSERT INTO repair_records(asset_id,fault_id,repair_date,technician,
                       action_taken,cost) VALUES(?,?,?,?,?,?)""",
                    (row["asset_id"], fault_id, repair_date, technician.strip(),
                     resolution.strip(), float(cost or 0)))
        open_faults = con.execute("""SELECT COUNT(*) FROM fault_reports
                                    WHERE asset_id=? AND status!='Resolved'""",
                                  (row["asset_id"],)).fetchone()[0]
        status = "Faulty" if open_faults else "Working"
        con.execute("UPDATE assets SET status=?, updated_at=? WHERE id=?",
                    (status, _now(), row["asset_id"]))

def add_movement(asset_code, movement_date, department_id, group_id, reason, moved_by):
    with connect() as con:
        asset = con.execute("""
            SELECT a.*, l.block, l.floor, l.room
            FROM assets a
            LEFT JOIN locations l ON l.id=a.location_id
            WHERE a.asset_code=?
        """, (asset_code,)).fetchone()
        if not asset:
            raise ValueError("Asset ID was not found.")
        destination = con.execute("""
            SELECT d.name AS department_name, g.parent_code, g.group_name,
                   g.location_id,
                   l.block, l.floor, l.room
            FROM asset_groups g
            LEFT JOIN departments d ON d.id=g.department_id
            LEFT JOIN locations l ON l.id=g.location_id
            WHERE g.id=? AND (? IS NULL OR g.department_id=?)
        """, (group_id, department_id, department_id)).fetchone()
        if not destination:
            raise ValueError("The selected destination Parent Group was not found.")
        old = " / ".join(x for x in [asset["block"], asset["floor"], asset["room"]] if x)
        new = " / ".join(x for x in [
            destination["department_name"],
            destination["parent_code"],
            destination["block"],
            destination["floor"],
            destination["room"],
        ] if x)
        con.execute("""INSERT INTO asset_movements(asset_id,movement_date,from_location,to_location,
                       reason,moved_by,department_id,group_id) VALUES(?,?,?,?,?,?,?,?)""",
                    (asset["id"], movement_date, old, new, reason.strip(), moved_by.strip(),
                     department_id, group_id))
        con.execute("""UPDATE assets SET department_id=?, group_id=?, location_id=?,
                       updated_at=? WHERE id=?""",
                    (department_id, group_id, destination["location_id"] if "location_id" in destination.keys() else None,
                     _now(), asset["id"]))
        return asset["asset_code"]

def dashboard_counts():
    with connect() as con:
        total = con.execute("SELECT COUNT(*) FROM assets").fetchone()[0]
        working = con.execute("SELECT COUNT(*) FROM assets WHERE status='Working'").fetchone()[0]
        faulty = con.execute("SELECT COUNT(*) FROM assets WHERE status='Faulty'").fetchone()[0]
        repair = con.execute("SELECT COUNT(*) FROM assets WHERE status='Under Repair'").fetchone()[0]
        groups = con.execute("SELECT COUNT(*) FROM asset_groups").fetchone()[0]
        open_faults = con.execute("SELECT COUNT(*) FROM fault_reports WHERE status!='Resolved'").fetchone()[0]
        return {"total":total,"working":working,"faulty":faulty,"repair":repair,
                "groups":groups,"open_faults":open_faults}

def clear_all_data():
    """Remove every user-created record while keeping the database schema."""
    tables = (
        "setup_assets",
        "asset_movements",
        "repair_records",
        "fault_reports",
        "assets",
        "computer_setups",
        "asset_groups",
        "locations",
        "departments",
    )
    con = connect()
    try:
        for table in tables:
            con.execute(f"DELETE FROM {table}")
        con.execute("DELETE FROM sqlite_sequence")
        con.commit()
    finally:
        con.close()

def get_asset_by_code(code):
    with connect() as con:
        return con.execute("SELECT * FROM assets WHERE asset_code=?", (code,)).fetchone()

def get_asset_details_by_code(code):
    with connect() as con:
        return con.execute("""
            SELECT a.*, d.name AS department_name, g.parent_code, g.group_name,
                   l.block, l.floor, l.room
            FROM assets a
            LEFT JOIN departments d ON d.id=a.department_id
            LEFT JOIN asset_groups g ON g.id=a.group_id
            LEFT JOIN locations l ON l.id=a.location_id
            WHERE a.asset_code=?
        """, (code.strip(),)).fetchone()

def export_assets_xlsx(path):
    from openpyxl import Workbook
    rows = list_assets("")
    wb = Workbook()
    ws = wb.active
    ws.title = "Assets"
    headers = ["Asset ID","Device Type","Brand","Model","Manufacturer Serial","Department",
               "Block","Floor","Room","Parent ID","Group","Setup ID","Setup","Staff",
               "Status","Processor","RAM","Storage","OS","Purchase Date","Warranty Expiry","Notes"]
    ws.append(headers)
    for a in rows:
        ws.append([a["asset_code"],a["device_type"],a["brand"],a["model"],a["manufacturer_serial"],
                   a["department_name"],a["block"],a["floor"],a["room"],a["parent_code"],
                   a["group_name"],a["setup_code"],a["setup_name"],a["staff_name"],a["status"],
                   a["processor"],a["ram"],a["storage"],a["operating_system"],a["purchase_date"],
                   a["warranty_expiry"],a["notes"]])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col[:100])
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len+2, 12), 32)
    wb.save(path)

def backup_database(path):
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    source = connect()
    dest = sqlite3.connect(path)
    try:
        source.backup(dest)
    finally:
        dest.close()
        source.close()
